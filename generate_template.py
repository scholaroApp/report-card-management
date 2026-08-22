"""
Generates template.xlsx - the data-entry spreadsheet the school office fills in.
One row per student. Column order matches the report card layout exactly.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SUBJECTS = ["English", "Hindi", "Mathematics", "Science", "Social Science", "Computer"]
SUBJECT_FIELDS = [
    ("P1", "Periodic Test 1 (/80)"),
    ("A", "A - PT Avg (/5)"),
    ("P2", "Periodic Test 2 (/80)"),
    ("B", "B - Multiple Assessment (/5)"),
    ("C", "C - Notebook (/5)"),
    ("D", "D - Subject Enrichment (/5)"),
    ("AnnualExam", "Annual Exam (/80)"),
]

CO_SCHOLASTIC = [
    "Courteousness", "Confidence", "CareOfBelongings", "Neatness",
    "RegularityPunctuality", "RespectForOthersBelongings",
    "DisciplineBehaviour", "AssignmentProjects",
]
HEALTH_PHYSICAL = ["WorkEducation", "ArtEducation", "HealthPhysicalEducation"]

STUDENT_INFO = [
    "Class", "Session", "AdmissionNo", "Name", "FatherName", "MotherName",
    "DOB", "PhoneNo", "Address", "Height", "Weight", "BloodGroup",
    "Vision", "Teeth", "OralHygiene",
]

def build_columns():
    cols = list(STUDENT_INFO)
    for subj in SUBJECTS:
        for field, _ in SUBJECT_FIELDS:
            cols.append(f"{subj.replace(' ', '')}_{field}")
    for act in CO_SCHOLASTIC:
        cols.append(f"{act}_T1")
        cols.append(f"{act}_T2")
    for act in HEALTH_PHYSICAL:
        cols.append(f"{act}_T1")
        cols.append(f"{act}_T2")
    cols += ["Attendance_T1", "Attendance_T2", "Remark_T1", "Remark_T2"]
    return cols

def sample_rows():
    row1 = {
        "Class": "IX", "Session": "2026-2027", "AdmissionNo": "PSM2026101",
        "Name": "Aarav Sharma", "FatherName": "Rajesh Sharma", "MotherName": "Sunita Sharma",
        "DOB": "12-04-2011", "PhoneNo": "9876543210",
        "Address": "H.No. 45, Sector 12, New Delhi",
        "Height": "152 cm", "Weight": "42 kg", "BloodGroup": "B+",
        "Vision": "6/6", "Teeth": "Good", "OralHygiene": "Good",
    }
    marks = {
        "English": [70, 4, 72, 4, 4, 4, 68], "Hindi": [65, 4, 68, 4, 4, 4, 70],
        "Mathematics": [75, 5, 78, 5, 4, 5, 74], "Science": [72, 4, 70, 4, 4, 4, 71],
        "Social Science": [68, 4, 66, 4, 4, 4, 65], "Computer": [78, 5, 80, 5, 5, 5, 76],
    }
    for subj, vals in marks.items():
        key = subj.replace(" ", "")
        for (field, _), v in zip(SUBJECT_FIELDS, vals):
            row1[f"{key}_{field}"] = v
    for act in CO_SCHOLASTIC:
        row1[f"{act}_T1"] = "A"
        row1[f"{act}_T2"] = "A"
    for act in HEALTH_PHYSICAL:
        row1[f"{act}_T1"] = "A"
        row1[f"{act}_T2"] = "A"
    row1["Attendance_T1"] = "88/90 days"
    row1["Attendance_T2"] = "85/90 days"
    row1["Remark_T1"] = "Sincere and hardworking. Keep it up."
    row1["Remark_T2"] = "Shows consistent improvement across subjects."
    return [row1]

def main():
    cols = build_columns()
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Students"
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c, name in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for r, row in enumerate(sample_rows(), 2):
        for c, name in enumerate(cols, 1):
            cell = ws.cell(row=r, column=c, value=row.get(name, ""))
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Instructions sheet
    ws2 = wb.create_sheet("Instructions")
    lines = [
        ("How to use this file", True),
        ("1. Fill one row per student in the 'Students' sheet. Do not rename or reorder columns.", False),
        ("2. Class must be exactly 'IX' or 'X'.", False),
        ("3. For each subject, fill P1, A, P2, B, C, D and AnnualExam (marks as per your internal scaling).", False),
        ("   A+B+C+D and Total are calculated automatically by the report card tool - do not add them here.", False),
        ("4. Co-Scholastic and Health & Physical Education grades: use A / B / C / D / E for each term.", False),
        ("5. Grade for each scholastic subject is calculated automatically from Total marks.", False),
        ("6. Save this file, then open index.html and click 'Import Excel' to load it - fully offline, no internet needed.", False),
        ("", False),
        ("Grading scale (for reference)", True),
        ("Scholastic: A1 91-100 | A2 81-90 | B1 71-80 | B2 61-70 | C1 51-60 | C2 41-50 | D 33-40 | E(Fail) 32 & below", False),
        ("Co-Scholastic: A=5 | B=4 | C=3 | D=2 | E=1", False),
    ]
    for r, (text, bold) in enumerate(lines, 1):
        cell = ws2.cell(row=r, column=1, value=text)
        cell.font = Font(bold=bold, size=12 if bold else 11)
    ws2.column_dimensions["A"].width = 110

    wb.save("template.xlsx")
    print("Saved template.xlsx with", len(cols), "columns")

if __name__ == "__main__":
    main()
